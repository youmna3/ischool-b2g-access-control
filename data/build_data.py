import openpyxl, json, re, collections

P1 = '/root/.claude/uploads/44f335b3-a42e-5e19-a0c5-65ee92abf128/7a100580-b2g_role_permissions_2026-09-21.xlsx'
P2 = '/root/.claude/uploads/44f335b3-a42e-5e19-a0c5-65ee92abf128/69e1bf17-b2g_user_access_audit_2026-09-21_1.xlsx'

wb1 = openpyxl.load_workbook(P1, data_only=True)
wb2 = openpyxl.load_workbook(P2, data_only=True)

# ---------- Roles ----------
roles_ws = wb1['Roles']
roles_rows = list(roles_ws.iter_rows(min_row=2, values_only=True))
roles = {}
for r in roles_rows:
    rid = str(r[0])
    roles[rid] = {
        'id': rid,
        'name': r[1],
        'scope': r[2],
        'boundOrg': r[3],
        'permScopeCount': r[4],
        'readOnlyScopes': r[5],
        'writeScopes': r[6],
        'permissionsRaw': r[7],
        'orgsAssigned': [o.strip() for o in (r[8] or '').split('|') if o.strip()] if r[8] else [],
        'b2gAssignments': r[9],
        'adminsInB2G': r[10],
        'activeAdminsInB2G': r[11],
        'adminsTotal': r[12],
        'primarySystemRoleOf': r[13],
        'createdAt': str(r[14]) if r[14] else None,
        'updatedAt': str(r[15]) if r[15] else None,
        'permissions': {},  # module -> {read, write, other: []}
    }

# ---------- Role x Permission ----------
rp_ws = wb1['Role × Permission']
rp_rows = list(rp_ws.iter_rows(min_row=2, values_only=True))
for r in rp_rows:
    rid = str(r[0])
    module = r[3]
    read = (r[4] == 'Yes')
    write = (r[5] == 'Yes')
    other = [a.strip() for a in (r[6] or '').split(',') if a.strip()]
    if rid in roles:
        roles[rid]['permissions'][module] = {'read': read, 'write': write, 'other': other}

# name -> canonical id map (first occurrence wins; note Super Admin ambiguity)
name_to_id = {}
for rid, rl in roles.items():
    nm = rl['name']
    if nm not in name_to_id:
        name_to_id[nm] = rid

roles_list = list(roles.values())
print('roles:', len(roles_list))
dupe_names = [n for n, c in collections.Counter([r['name'] for r in roles_list]).items() if c > 1]
print('duplicate role names:', dupe_names)

# ---------- Users (Summary) ----------
u_ws = wb2['Summary']
u_rows = list(u_ws.iter_rows(min_row=2, values_only=True))
users = []
missing_role_names = collections.Counter()

def parse_org_role_pairs(s):
    pairs = []
    if not s or s == '(none)':
        return pairs
    for part in s.split('|'):
        part = part.strip()
        if '→' in part:
            org, role = part.split('→', 1)
            pairs.append((org.strip(), role.strip()))
    return pairs

for r in u_rows:
    uid = r[0]
    if uid is None:
        continue
    uid = str(int(uid)) if isinstance(uid, float) else str(uid)
    name = r[1]
    email = r[2]
    primary_role = r[3]
    subrole = r[4]
    orgs_str = r[5]
    org_role_pairs_str = r[6]
    status = r[7]
    department = r[8]
    user_type = r[9]
    team_lead = r[10]
    dept_head = r[11]
    home_org = r[12]
    org_count = r[14]
    why_in_scope = r[15]
    last_activity = str(r[20]) if r[20] else None
    created_at = str(r[21]) if r[21] else None

    pairs = parse_org_role_pairs(org_role_pairs_str)
    role_assignments = []
    for org, role_name in pairs:
        rid = name_to_id.get(role_name)
        if rid is None:
            missing_role_names[role_name] += 1
        role_assignments.append({
            'org': org,
            'roleName': role_name,
            'roleId': rid,
            'status': 'active',
            'addedAt': created_at,
            'addedBy': 'system-import',
            'removedAt': None,
            'removedBy': None,
        })

    users.append({
        'id': uid,
        'name': name,
        'email': email,
        'primaryRole': primary_role,
        'subRole': subrole,
        'status': status,
        'department': department if department and department != '(none)' else 'Unassigned',
        'userType': user_type,
        'teamLead': (team_lead == 'Yes'),
        'deptHead': (dept_head == 'Yes'),
        'homeOrg': home_org if home_org and home_org != '(none)' else None,
        'orgCount': org_count,
        'whyInScope': why_in_scope,
        'lastActivity': last_activity,
        'createdAt': created_at,
        'roles': role_assignments,
        'deleted': False,
        'deletedAt': None,
        'deletedBy': None,
    })

print('users:', len(users))
print('missing role name matches (top 15):', missing_role_names.most_common(15))

depts = collections.Counter([u['department'] for u in users])
print('departments:', depts)

with open('roles.json', 'w') as f:
    json.dump(roles_list, f, ensure_ascii=False)
with open('users.json', 'w') as f:
    json.dump(users, f, ensure_ascii=False)

print('done')
