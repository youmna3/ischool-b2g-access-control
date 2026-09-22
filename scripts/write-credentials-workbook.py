import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


payload = json.load(sys.stdin)
output_path = Path(sys.argv[1]).resolve()
output_path.parent.mkdir(parents=True, exist_ok=True)

workbook = Workbook()
workbook.remove(workbook.active)


def add_sheet(title, rows, columns):
    sheet = workbook.create_sheet(title)
    sheet.append([label for _, label in columns])
    for row in rows:
        sheet.append([row.get(key, "") for key, _ in columns])

    header_fill = PatternFill("solid", fgColor="056FEC")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, (key, label) in enumerate(columns, 1):
        values = [str(row.get(key, "")) for row in rows]
        width = min(55, max([len(label), *(len(value) for value in values)] if values else [len(label)]) + 2)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
        if key in {"email", "temporary_password", "auth_user_id", "source_ids", "source_id"}:
            for cell in sheet[get_column_letter(column_index)][1:]:
                cell.number_format = "@"
    return sheet


instructions = workbook.create_sheet("Instructions")
instructions.append(["CONFIDENTIAL — Temporary Account Credentials"])
instructions["A1"].font = Font(bold=True, color="FFFFFF", size=14)
instructions["A1"].fill = PatternFill("solid", fgColor="DE1F1F")
instructions.append(["Generated at", payload["generated_at"]])
instructions.append(["Handling", "Store and transfer this workbook securely. Delete it after credentials are distributed."])
instructions.append(["Passwords", "Each unique email has a unique temporary password. Ask every user to change it after first login."])
instructions.append(["Duplicates", "Duplicate seed emails share one Supabase Auth identity and therefore one password."])
instructions.column_dimensions["A"].width = 22
instructions.column_dimensions["B"].width = 100

add_sheet("Credentials", payload["credentials"], [
    ("source_ids", "Seed User ID(s)"),
    ("seed_record_count", "Seed Records"),
    ("name", "Name"),
    ("email", "Email"),
    ("temporary_password", "Temporary Password"),
    ("auth_user_id", "Supabase Auth User ID"),
    ("auth_action", "Auth Action"),
    ("role", "Application Role"),
    ("department", "Account Department"),
    ("active", "Active"),
    ("status", "Status"),
])

add_sheet("Duplicate Seed Emails", payload["duplicates"], [
    ("source_id", "Seed User ID"),
    ("name", "Name"),
    ("email", "Email"),
    ("department", "Department"),
    ("duplicate_email", "Shared Auth Email"),
])

add_sheet("Invalid Seed Rows", payload["invalid"], [
    ("source_id", "Seed User ID"),
    ("name", "Name"),
    ("email", "Email"),
    ("department", "Department"),
    ("reason", "Reason"),
])

workbook.save(output_path)
